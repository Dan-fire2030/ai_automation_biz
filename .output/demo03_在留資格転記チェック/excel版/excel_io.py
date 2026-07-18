"""Excel入出力（openpyxl）。公式様式xlsxのコピー・転記・読み取り・チェック結果の書き出し。

セル位置の根拠は ../00_セルマッピング.md 参照。
"""

import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from config import CHECK_HEADERS, CONFIG, FORM_CELLS, VERDICT_COLORS
from pure import build_application_values


def create_form_copy(output_path):
    """公式様式xlsxをコピーして新しいファイルを作り、Workbookとして開いて返す。"""
    template = Path(CONFIG["TEMPLATE_PATH"])
    if not template.exists():
        raise RuntimeError(f"公式様式が見つかりません: {template}")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, output_path)
    return openpyxl.load_workbook(output_path)


def get_form_sheet(workbook):
    """様式シート（申請人等作成用1）を取得する。無ければ明確なエラー。"""
    name = CONFIG["FORM_SHEET_NAME"]
    if name not in workbook.sheetnames:
        raise RuntimeError(
            f"シート「{name}」が見つかりません。公式様式のxlsxか確認してください（シート: {workbook.sheetnames}）"
        )
    return workbook[name]


def fill_form_sheet(sheet, values):
    """申請書シートに値を転記する。空の値（None/空文字）は書き込まない。"""
    cell_value_pairs = (
        (FORM_CELLS["NATIONALITY"], values.get("nationality")),
        (FORM_CELLS["BIRTH_YEAR"], values.get("birth_year")),
        (FORM_CELLS["BIRTH_MONTH"], values.get("birth_month")),
        (FORM_CELLS["BIRTH_DAY"], values.get("birth_day")),
        (FORM_CELLS["NAME"], values.get("name")),
        (FORM_CELLS["PASSPORT_NUMBER"], values.get("passport_number")),
        (FORM_CELLS["PASSPORT_EXP_YEAR"], values.get("passport_exp_year")),
        (FORM_CELLS["PASSPORT_EXP_MONTH"], values.get("passport_exp_month")),
        (FORM_CELLS["PASSPORT_EXP_DAY"], values.get("passport_exp_day")),
        (FORM_CELLS["STATUS_OF_RESIDENCE"], values.get("status_of_residence")),
        (FORM_CELLS["PERIOD_OF_STAY"], values.get("period_of_stay")),
        (FORM_CELLS["STAY_EXP_YEAR"], values.get("stay_exp_year")),
        (FORM_CELLS["STAY_EXP_MONTH"], values.get("stay_exp_month")),
        (FORM_CELLS["STAY_EXP_DAY"], values.get("stay_exp_day")),
        (FORM_CELLS["CARD_NUMBER"], values.get("card_number")),
    )
    for a1, value in cell_value_pairs:
        if value is None or value == "":
            continue
        sheet[a1] = value

    # 性別は○囲み式のため、該当セルのハイライトで代替する
    highlight = PatternFill(
        start_color=CONFIG["SEX_HIGHLIGHT_COLOR"], end_color=CONFIG["SEX_HIGHLIGHT_COLOR"], fill_type="solid"
    )
    if values.get("sex") == "男":
        sheet[FORM_CELLS["SEX_MALE"]].fill = highlight
    elif values.get("sex") == "女":
        sheet[FORM_CELLS["SEX_FEMALE"]].fill = highlight


def read_form_sheet(sheet):
    """申請書シートから突合対象の記載を読み取り、チェック用の項目リストを返す。"""
    def cell(a1):
        return sheet[a1].value

    raw = {
        "nationality": cell(FORM_CELLS["NATIONALITY"]),
        "birth_year": cell(FORM_CELLS["BIRTH_YEAR"]),
        "birth_month": cell(FORM_CELLS["BIRTH_MONTH"]),
        "birth_day": cell(FORM_CELLS["BIRTH_DAY"]),
        "name": cell(FORM_CELLS["NAME"]),
        "sex": read_sex_selection(sheet),
        "passport_number": cell(FORM_CELLS["PASSPORT_NUMBER"]),
        "passport_exp_year": cell(FORM_CELLS["PASSPORT_EXP_YEAR"]),
        "passport_exp_month": cell(FORM_CELLS["PASSPORT_EXP_MONTH"]),
        "passport_exp_day": cell(FORM_CELLS["PASSPORT_EXP_DAY"]),
        "status_of_residence": cell(FORM_CELLS["STATUS_OF_RESIDENCE"]),
        "period_of_stay": cell(FORM_CELLS["PERIOD_OF_STAY"]),
        "stay_exp_year": cell(FORM_CELLS["STAY_EXP_YEAR"]),
        "stay_exp_month": cell(FORM_CELLS["STAY_EXP_MONTH"]),
        "stay_exp_day": cell(FORM_CELLS["STAY_EXP_DAY"]),
        "card_number": cell(FORM_CELLS["CARD_NUMBER"]),
    }
    return build_application_values(raw)


def read_sex_selection(sheet):
    """性別欄のハイライト（○囲みの代替）からどちらが選択されているかを読む。"""
    highlight = CONFIG["SEX_HIGHLIGHT_COLOR"].upper()

    def is_highlighted(a1):
        fill = sheet[a1].fill
        rgb = getattr(fill.start_color, "rgb", "") or ""
        return fill.fill_type == "solid" and str(rgb).upper().endswith(highlight)

    if is_highlighted(FORM_CELLS["SEX_MALE"]):
        return "男"
    if is_highlighted(FORM_CELLS["SEX_FEMALE"]):
        return "女"
    return ""


def write_check_results(workbook, check):
    """チェック結果シートを（作り直して）書き出す。同じxlsxにシートが増える形。"""
    name = CONFIG["CHECK_SHEET_NAME"]
    if name in workbook.sheetnames:
        del workbook[name]
    sheet = workbook.create_sheet(name, 0)

    bold = Font(bold=True)
    for col, header in enumerate(CHECK_HEADERS, start=1):
        header_cell = sheet.cell(row=1, column=col, value=header)
        header_cell.font = bold

    for row, result in enumerate(check["results"], start=2):
        row_values = (
            result["item"],
            result["source"],
            result["application_value"],
            result["original_value"],
            result["verdict"],
            result["comment"],
        )
        color = VERDICT_COLORS.get(result["verdict"])
        for col, value in enumerate(row_values, start=1):
            data_cell = sheet.cell(row=row, column=col, value=value)
            if color:
                data_cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    summary_row = len(check["results"]) + 3
    sheet.cell(row=summary_row, column=1, value="まとめ: " + check.get("summary", ""))

    widths = {"A": 18, "B": 12, "C": 26, "D": 26, "E": 8, "F": 46}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
